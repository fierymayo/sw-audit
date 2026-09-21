import os

import config

XLSX_HEADERS = ["ID", "Handle", "Title", "Count Now", "Suggested Rule", "Adds",
                "Rule type", "Count After", "Approve", "Note"]
_WIDTHS = [16, 30, 42, 11, 60, 8, 16, 12, 10, 24]


def _split(rows):
    ready = [r for r in rows if r["classification"] in ("AUTOMATABLE", "COLOUR_VERIFY")]
    review = [r for r in rows if r["classification"] == "REVIEW"]
    return ready, review


def _coll_id(row):
    return row["admin_url"].rsplit("/", 1)[-1]


def build_xlsx(rows, stamp, log=print):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        log("  XLSX skipped: openpyxl is not installed (pip install openpyxl).")
        return None

    ready, review = _split(rows)
    header_fill = PatternFill("solid", fgColor="1A1A2E")
    header_font = Font(color="FFFFFF", bold=True)
    approve_fill = PatternFill("solid", fgColor="FFF2CC")
    link_font = Font(color="0563C1", underline="single")

    wb = Workbook()

    def sheet(ws, data):
        ws.append(XLSX_HEADERS)
        for i, w in enumerate(_WIDTHS, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for cell in ws[1]:
            cell.fill, cell.font = header_fill, header_font
        for r in data:
            cid = _coll_id(r)
            ws.append([cid, r["handle"], r["title"], r["count"], r["suggested_rule"],
                       r["adds"], r["rule_type"], r["count_after"], "", ""])
            row = ws.max_row
            id_cell = ws.cell(row=row, column=1)
            id_cell.number_format = "@"
            id_cell.hyperlink = r["admin_url"]
            id_cell.font = link_font
            handle_cell = ws.cell(row=row, column=2)
            handle_cell.hyperlink = f"{config.STOREFRONT}/collections/{r['handle']}"
            handle_cell.font = link_font
            ws.cell(row=row, column=9).fill = approve_fill
        ws.freeze_panes = "A2"

    ws1 = wb.active
    ws1.title = "Ready to Automate"
    sheet(ws1, ready)
    sheet(wb.create_sheet("Needs Review"), review)

    path = os.path.join(config.run_dir(stamp), f"Collections-to-Automate-{stamp}.xlsx")
    os.makedirs(config.OUT_DIR, exist_ok=True)
    wb.save(path)
    log(f"  wrote {path}  (ready={len(ready)}, review={len(review)})")
    return path


def build_pdf(rows, stamp, log=print):
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import landscape, letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError:
        log("  collections PDF skipped: reportlab is not installed.")
        return None

    ready, review = _split(rows)
    styles = getSampleStyleSheet()
    body = styles["BodyText"]
    body.fontSize = 7.5
    story = [Paragraph("Soccer Wearhouse — Collections to Automate", styles["Title"]),
             Paragraph("Additive rules only: every suggestion keeps current members and can "
                       "only ADD products. Approve in the matching spreadsheet.", styles["Italic"])]

    def section(title, data):
        if not data:
            return
        story.append(Spacer(1, 12))
        story.append(Paragraph(title, styles["Heading2"]))
        table_rows = [["Collection", "Now", "Suggested rule", "Adds", "After"]]
        for r in data:
            table_rows.append([
                Paragraph(f'<link href="{r["admin_url"]}">{r["title"]}</link>', body),
                str(r["count"]), Paragraph(r["suggested_rule"], body),
                str(r["adds"]), str(r["count_after"])])
        t = Table(table_rows, colWidths=[2.6 * inch, 0.5 * inch, 4.6 * inch, 0.5 * inch, 0.6 * inch],
                  hAlign="LEFT", repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a1a2e")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f2f7")]),
        ]))
        story.append(t)

    section(f"Ready to automate ({len(ready)})", ready)
    section(f"Needs review ({len(review)})", review)

    path = os.path.join(config.run_dir(stamp), f"Collections-to-Automate-{stamp}.pdf")
    SimpleDocTemplate(path, pagesize=landscape(letter),
                      title="SW Collections to Automate").build(story)
    log(f"  wrote {path}")
    return path