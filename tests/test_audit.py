import datetime
import json
import os
import sys
import tempfile
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import collections_audit as ca
import config
import report
from catalog import load_products
from matcher import compile_rules, match_title, normalize, strip_fc
from passes import (pass_addon_coverage, pass_filled_check, pass_fill_rates, pass_filter_blanks,
                    pass_forecast, pass_vendor_audit)
from rules import config_lint, load_task_rules
from subcat_sim import classification_type, classify, pass_subcat_sim

_SIM_REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIXTURES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "fixtures")
NOW = datetime.datetime(2026, 9, 21, 7, 0, tzinfo=datetime.timezone.utc)


def load_fixture_state():
    cfg = load_task_rules(os.path.join(FIXTURES, "test-config.json"))
    products = load_products(os.path.join(FIXTURES, "sample.jsonl"))
    return cfg, products


class TestMatcher(unittest.TestCase):
    def test_normalize_parity(self):
        self.assertEqual(normalize("Lautaro Martínez"), "lautaro martinez")
        self.assertEqual(normalize("Müller / Straße"), "muller strasse")
        self.assertEqual(normalize("Jerseys & Gear"), "jerseys and gear")
        self.assertEqual(normalize("St. Louis City-SC (Home)"), "st louis city sc home")

    def test_longest_wins(self):
        compiled = compile_rules([
            {"canonical_value": "Martinez", "safe_keywords": ["lautaro martinez"]},
            {"canonical_value": "Inter Miami CF", "safe_keywords": ["inter"]},
        ])
        best, ambiguous = match_title("Inter Lautaro Martinez Jersey", compiled)
        self.assertEqual(best, "Martinez")
        self.assertFalse(ambiguous)

    def test_equal_length_tie_is_ambiguous(self):
        compiled = compile_rules([
            {"canonical_value": "TieA", "safe_keywords": ["zzzz"]},
            {"canonical_value": "TieB", "safe_keywords": ["yyyy"]},
        ])
        best, ambiguous = match_title("Legends zzzz yyyy Scarf", compiled)
        self.assertIsNone(best)
        self.assertTrue(ambiguous)

    def test_padded_word_boundaries(self):
        compiled = compile_rules([{"canonical_value": "Son", "safe_keywords": ["son"]}])
        self.assertEqual(match_title("Season Opener Jersey", compiled), (None, False))
        self.assertEqual(match_title("Son Heung Min Jersey", compiled)[0], "Son")

    def test_strip_fc(self):
        self.assertEqual(strip_fc("leon fc"), "leon")
        self.assertEqual(strip_fc("villarreal cf"), "villarreal")


class TestReasonEngine(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cfg, cls.products = load_fixture_state()
        cls.blanks = pass_filter_blanks(cls.products, cls.cfg, lag_hours=48, now=NOW)
        cls.by_handle = {r["handle"]: r for r in cls.blanks["player_filter"]}
        cls.club_by_handle = {r["handle"]: r for r in cls.blanks["club_filter"]}

    def reason(self, handle):
        return self.by_handle[handle]["reason"]

    def test_sync_miss_vs_lag(self):
        self.assertEqual(self.reason("p1"), "SYNC_MISS")
        self.assertEqual(self.reason("p2"), "SYNC_LAG")

    def test_keyword_gap_on_bare_surname(self):
        self.assertEqual(self.reason("p3"), "KEYWORD_GAP")
        self.assertEqual(self.by_handle["p3"]["matched_value"], "Davies")

    def test_accent_fold_matches(self):
        self.assertEqual(self.reason("p4"), "SYNC_MISS")
        self.assertEqual(self.by_handle["p4"]["matched_value"], "Martinez")

    def test_no_signal_not_flagged(self):
        self.assertNotIn("p5", self.by_handle)
        self.assertNotIn("p5", self.club_by_handle)

    def test_ambiguous(self):
        self.assertEqual(self.reason("p6"), "AMBIGUOUS")

    def test_invalid_choice(self):
        self.assertEqual(self.club_by_handle["p7"]["reason"], "INVALID_CHOICE")

    def test_no_rule_value_from_catalog_vocab(self):
        self.assertEqual(self.club_by_handle["p11"]["reason"], "NO_RULE_VALUE")
        self.assertEqual(self.club_by_handle["p11"]["matched_value"], "Tigres UANL")

    def test_inactive_ignored(self):
        for rows in self.blanks.values():
            self.assertNotIn("p10", {r["handle"] for r in rows})


class TestOtherPasses(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cfg, cls.products = load_fixture_state()

    def test_fill_rates_and_diff(self):
        s = pass_fill_rates(self.products, self.cfg)
        self.assertEqual(s["active"], 14)
        self.assertEqual(s["fill_club_filter"]["filled"], 3)
        self.assertEqual(s["fill_sibling_products"]["filled"], 2)
        self.assertEqual(s["type_misc_active"], 1)
        prev = dict(s, fill_club_filter={"filled": 2, "pct": 8.3}, active=13)
        s2 = pass_fill_rates(self.products, self.cfg, prev=prev, prev_name="prev.json")
        self.assertEqual(s2["delta_vs_previous"]["fill_club_filter"]["filled"], 1)
        self.assertEqual(s2["delta_vs_previous"]["active"], 1)

    def test_sibling_precision_counts_multi_member_groups_only(self):
        # group_x has p13 (blank) + p15 (filled); group_solo and group_lonely have one member each
        scope = pass_fill_rates(self.products, self.cfg)["sibling_scope"]
        self.assertEqual(scope, {"products_with_group_tag": 3, "multi_member_groups": 1,
                                 "products_in_multi_member_groups": 2,
                                 "filled_in_multi_member_groups": 1})

    def test_addon_coverage(self):
        cov = pass_addon_coverage(self.products, self.cfg)
        self.assertEqual([r["handle"] for r in cov["missing_tag"]], ["p6"])
        self.assertEqual([r["handle"] for r in cov["stale_tag"]], ["p9"])

    def test_vendor_audit(self):
        v = pass_vendor_audit(self.products, self.cfg)
        findings = {r["vendor"]: r for r in v["rows"]}
        self.assertEqual(findings["NIKE"]["finding"], "NOT_NORMALIZED")
        self.assertEqual(findings["NIKE"]["suggested"], "Nike")
        self.assertEqual(findings["Kwik-Goal"]["finding"], "CASING_ALIAS_CANDIDATE")
        self.assertEqual(findings["Kwik-Goal"]["suggested"], "Kwik Goal")
        self.assertEqual(findings["soccerwearhouse.com"]["finding"], "NOT_NORMALIZED")
        self.assertEqual(findings["soccerwearhouse.com"]["suggested"], "Soccer Wearhouse")

    def test_lint_flags_invalid_choice_config(self):
        findings = config_lint(self.cfg)
        errors = [f for f in findings if f["level"] == "ERROR"]
        self.assertTrue(any("Leon FC" in f["message"] for f in errors))


class TestFilledCheck(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cfg, cls.products = load_fixture_state()
        cls.result = pass_filled_check(cls.products, cls.cfg)
        cls.club = {r["handle"]: r for r in cls.result["rows"]["club_filter"]}

    def test_orphan_rows_and_counts(self):
        row = self.club["p12"]
        self.assertEqual((row["finding"], row["stored_value"], row["predicted_value"]),
                         ("FILLED_ORPHAN", "Tigres UANL", ""))
        self.assertEqual(self.result["orphans"]["club_filter"], {"Tigres UANL": 1})
        self.assertEqual(self.result["orphans"]["country_filter"], {"Wales": 1})
        self.assertEqual(self.result["orphans"]["player_filter"], {})

    def test_divergent(self):
        row = self.club["p14"]
        self.assertEqual((row["finding"], row["stored_value"], row["predicted_value"]),
                         ("FILLED_DIVERGENT", "Inter Miami CF", "FC Barcelona"))

    def test_agreeing_and_unpredicted_values_are_not_flagged(self):
        self.assertNotIn("p1", self.club)
        self.assertEqual(self.result["rows"]["player_filter"], [])

    def test_skipped_in_fallback_mode(self):
        self.assertIsNone(pass_filled_check(self.products, None))


class TestForecast(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cfg, cls.products = load_fixture_state()

    def forecast(self, rules):
        return pass_forecast(self.products, self.cfg, "player", rules)

    def test_delta_rule_fills_known_blank(self):
        out = self.forecast([{"canonical_value": "Jordan Davies", "safe_keywords": ["jordan davies"]}])
        row = out["rows"][0]
        self.assertEqual((row["canonical"], row["keyword_count"], row["blanks_filled"], row["flags"]),
                         ("Jordan Davies", 1, 1, ""))
        self.assertEqual(row["sample_title"], "Jordan Davies Wales Tee")
        self.assertEqual((out["new_rules"], out["blanks_filled"], out["blanks"]), (1, 1, 13))

    def test_existing_longer_keyword_still_steals(self):
        # "home" also hits p1, but existing "messi" is longer and wins there
        out = self.forecast([{"canonical_value": "Home Guy", "safe_keywords": ["home"]}])
        self.assertEqual(out["rows"][0]["blanks_filled"], 2)

    def test_collision_with_existing_keyword_is_flagged_and_ambiguous(self):
        out = self.forecast([{"canonical_value": "Lionel", "safe_keywords": ["messi"]}])
        row = out["rows"][0]
        self.assertIn("COLLISION_EXISTING(messi -> Messi)", row["flags"])
        self.assertEqual(row["blanks_filled"], 0)

    def test_duplicate_within_delta_is_flagged(self):
        out = self.forecast([{"canonical_value": "A1", "safe_keywords": ["zed"]},
                             {"canonical_value": "B1", "safe_keywords": ["zed"]}])
        flags = {r["canonical"]: r["flags"] for r in out["rows"]}
        self.assertEqual(flags["A1"], "DUPLICATE_IN_DELTA(zed also in B1)")
        self.assertEqual(flags["B1"], "DUPLICATE_IN_DELTA(zed also in A1)")

    def test_needs_live_rules(self):
        self.assertIsNone(pass_forecast(self.products, None, "player", [{"canonical_value": "X"}]))


class TestSnapshot(unittest.TestCase):
    def test_writes_md_with_deltas_only_when_previous_exists(self):
        cfg, products = load_fixture_state()
        summary = pass_fill_rates(products, cfg)
        summary["filled_orphans"] = pass_filled_check(products, cfg)["orphans"]
        prev = dict(summary, active=13, fill_club_filter={"filled": 2, "pct": 8.3})
        with tempfile.TemporaryDirectory() as td, mock.patch.object(config, "OUT_DIR", td):
            cache = os.path.join(td, "products.jsonl")
            open(cache, "w").close()
            path = report.write_snapshot_md(summary, None, cache, log=lambda *_: None)
            self.assertEqual(os.path.basename(path), f"catalog-snapshot-{summary['generated_at'][:10]}.md")
            with open(path, encoding="utf-8") as fh:
                first = fh.read()
            self.assertIn("| Active products | 14 |  |", first)
            self.assertIn("- club_filter: Tigres UANL ×1", first)
            self.assertIn("Sibling precision: 1 of 2 products in multi-member groups filled (50.0%)", first)
            self.assertIn(f"rules file {summary['rules_file_date']}", first)
            self.assertIn("cache products.jsonl", first)

            report.write_snapshot_md(summary, prev, cache, log=lambda *_: None)
            with open(path, encoding="utf-8") as fh:
                second = fh.read()
            self.assertIn("| Active products | 14 | +1 |", second)
            self.assertIn("| club_filter | 3 | ", second)


class TestRulesFileDate(unittest.TestCase):
    def test_prefers_generated_at(self):
        cfg = load_task_rules(os.path.join(FIXTURES, "test-config.json"))
        self.assertEqual(cfg["_meta"]["file_date"], "2026-09-21")

    def test_falls_back_to_mtime(self):
        with tempfile.TemporaryDirectory() as td:
            path = os.path.join(td, "cfg.json")
            with open(path, "w") as fh:
                json.dump({}, fh)
            ts = datetime.datetime(2025, 3, 4, 12).timestamp()
            os.utime(path, (ts, ts))
            self.assertEqual(load_task_rules(path)["_meta"]["file_date"], "2025-03-04")


class TestGenerator(unittest.TestCase):
    def test_parses_export_shape(self):
        import make_task_configs as gen
        export = {
            "id": "6af0e5de-4ae5-4163-9bed-cd120471b60a",
            "name": "Sync player",
            "options": {
                "metafield_key__string_required": "player_filter",
                "write_mode__string_required": "add_only",
                "keyword_rules_json__multiline_required":
                    json.dumps([{"canonical_value": "Messi", "safe_keywords": ["messi"]}]),
                "allowed_values_json__multiline": json.dumps(["Messi"]),
            },
        }
        with tempfile.TemporaryDirectory() as td:
            src = os.path.join(td, "player.json")
            out = os.path.join(td, "out.json")
            with open(src, "w") as fh:
                json.dump(export, fh)
            cfg = gen.build([src], out_path=out)
            self.assertEqual(len(cfg["player"]["keyword_rules"]), 1)
            self.assertEqual(cfg["player"]["write_mode"], "add_only")
            self.assertTrue(os.path.exists(out))



def _sim_cfg():
    cfg = load_task_rules(os.path.join(_SIM_REPO, "task-configs.json"))
    if not cfg or not (cfg.get("subcat") or {}).get("product_type_map"):
        raise unittest.SkipTest("real task-configs.json with subcat map required")
    return cfg


def _sim_p(n, title, ptype, tags, status="ACTIVE"):
    return {"handle": f"s{n}", "id": str(n), "title": title, "type": ptype,
            "tags": tags, "_tags_lc": [t.lower() for t in tags], "status": status}


class TestSubcatClassify(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cfg = _sim_cfg()
        cls.tmap = cls.cfg["subcat"]["product_type_map"]

    def ct(self, ptype):
        return classification_type(ptype, self.tmap)

    def test_type_mapping(self):
        self.assertEqual(self.ct("shoes"), "Footwear")
        self.assertEqual(self.ct("HATS"), "Hats")
        self.assertEqual(self.ct("unclassified"), "Misc")
        self.assertEqual(self.ct("Jerseys"), "Jerseys")

    def test_pre_wrap_asymmetry(self):
        title = "Mueller Pre Wrap Big Roll (Blue)"
        self.assertEqual(classify(title, self.ct("Accessories")), "SubCat_Accessory-Other")
        self.assertEqual(classify(title, self.ct("Misc")), "SubCat_Sock-Tape")

    def test_substring_not_word_boundary(self):
        self.assertEqual(classify("Nike Snapback Hat", "Hats"), "SubCat_Hat-Snapback")
        self.assertEqual(classify("Nike Gripknit Socks", "Socks"), "SubCat_Grip-Socks")

    def test_gk_uses_title_words(self):
        self.assertEqual(classify("Reusch GK-Shorts Pro", "Shorts"), "SubCat_GK-Shorts")
        self.assertIsNone(classify("Reusch GKX Shorts", "Shorts"))
        self.assertIsNone(classify("Nike Running Shorts", "Shorts"))

    def test_accessories_exclusions(self):
        self.assertIsNone(classify("Mexico Flag", "Accessories"))
        self.assertIsNone(classify("Vizari Referee Wallet", "Accessories"))
        self.assertEqual(classify("Chelsea Air Freshener", "Air Freshener"), "SubCat_Air-Freshener")

    def test_guarded_branches(self):
        self.assertIsNone(classify("New Era Tape Logo Cap", "Accessories"))
        self.assertIsNone(classify("Jordan Jumpman Pump Hat", "Accessories"))
        self.assertEqual(classify("Nike Essential Ball Pump", "Misc"), "SubCat_Ball-Pump")

    def test_misc_no_fallback(self):
        self.assertIsNone(classify("Random Novelty Item", "Misc"))

    def test_footwear_slides_only(self):
        self.assertEqual(classify("adidas Adilette Slides", "Footwear"), "SubCat_Slides")
        self.assertIsNone(classify("adidas Predator FG Cleats", "Footwear"))


class TestSubcatPass(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cfg = _sim_cfg()

    def run_pass(self, products):
        return pass_subcat_sim(products, self.cfg)

    def test_findings(self):
        out = self.run_pass([
            _sim_p(1, "FC Barcelona Scarf", "Scarves", []),
            _sim_p(2, "Chelsea Poster", "Posters", ["SubCat_Scarf"]),
            _sim_p(3, "Liverpool Scarf", "Scarves", ["SubCat_Scarf"]),
            _sim_p(4, "Arsenal Jersey", "Jerseys", ["SubCat_Collectible"]),
            _sim_p(5, "Multi Tag Scarf", "Scarves", ["SubCat_Scarf", "SubCat_Poster"]),
            _sim_p(6, "Draft Scarf", "Scarves", [], status="DRAFT"),
        ])
        by = {r["handle"]: r["finding"] for r in out["rows"]}
        self.assertEqual(by, {"s1": "MISSING_SUBCAT", "s2": "WILL_REPLACE",
                              "s4": "UNPREDICTED_EXISTING", "s5": "WILL_REPLACE"})
        self.assertEqual(out["counts"], {"MISSING_SUBCAT": 1, "WILL_REPLACE": 2,
                                         "UNPREDICTED_EXISTING": 1})

    def test_eligibility_columns(self):
        out = self.run_pass([_sim_p(7, "Inter Miami Scarf", "Scarves", ["SubCat_Bag-Other"])])
        row = out["rows"][0]
        self.assertTrue(row["predicted_eligible"])
        self.assertFalse(row["currently_eligible"])
        self.assertFalse(row["has_addon_tag"])

    def test_pre_wrap_pair_through_pass(self):
        out = self.run_pass([
            _sim_p(8, "Mueller Pre Wrap Big Roll", "Accessories", []),
            _sim_p(9, "Mueller Pre Wrap Big Roll", "Misc", []),
        ])
        pred = {r["handle"]: (r["predicted"], r["predicted_eligible"]) for r in out["rows"]}
        self.assertEqual(pred["s8"], ("SubCat_Accessory-Other", False))
        self.assertEqual(pred["s9"], ("SubCat_Sock-Tape", True))

    


def _c_p(n, title, ptype, **filters):
    p = {"gid": f"gid://shopify/Product/{n}", "id": str(n), "handle": f"c{n}",
         "title": title, "type": ptype, "status": "ACTIVE", "tags": [], "_tags_lc": [],
         "club_filter": "", "country_filter": "", "player_filter": "", "tournament_filter": ""}
    p.update(filters)
    return p


def _coll(n, title, count, has_rule=False):
    return {"gid": f"gid://shopify/Collection/{n}", "id": str(n), "title": title,
            "handle": f"col{n}", "count": count, "sources_n": 1 if has_rule else 0,
            "conditions": [], "has_rule": has_rule}


class TestCandidateEngine(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.cfg = _sim_cfg()
        active = ([_c_p(i, f"Poster {i}", "Posters") for i in range(1, 9)]
                  + [_c_p(20, "Real Madrid Home Jersey", "Jerseys", club_filter="Real Madrid"),
                     _c_p(21, "Real Madrid Away Jersey", "Jerseys", club_filter="Real Madrid"),
                     _c_p(22, "Real Madrid Beanie", "Hats", club_filter="Real Madrid"),
                     _c_p(23, "Blue Predator Cleats", "Footwear"),
                     _c_p(24, "White Mercurial Cleats", "Footwear")])
        cls.ctx = {"active": active,
                   "type_counts": {"Posters": 8, "Jerseys": 2, "Hats": 1, "Footwear": 2}}

    def classify_one(self, title, count=10, members=None):
        rows = ca.find_candidates([_coll(1, title, count)], self.cfg,
                                  products_ctx=self.ctx, members=members)
        return rows[0]

    def test_review_triggers(self):
        self.assertIn("PDP", self.classify_one("Jersey: Lionel Messi")["suggested_rule"])
        self.assertEqual(self.classify_one("Custom Nameset Builder")["classification"], "REVIEW")
        self.assertEqual(self.classify_one("Club Am\u00e9rica Jerseys")["classification"], "REVIEW")
        self.assertEqual(self.classify_one("St. Pauli Gear")["classification"], "REVIEW")
        self.assertEqual(self.classify_one("Deportivo de Cali")["classification"], "REVIEW")

    def test_brand_only_review(self):
        row = self.classify_one("adidas Collection")
        self.assertEqual(row["classification"], "REVIEW")
        self.assertIn("brand", row["suggested_rule"])

    def test_colour_verify_tier(self):
        row = self.classify_one("Blue Soccer Cleats")
        self.assertEqual(row["classification"], "COLOUR_VERIFY")
        self.assertEqual(row["rule_type"], "Title colour")
        self.assertIn("'blue'", row["suggested_rule"])

    def test_type_anchor_and_broad(self):
        anchored = self.classify_one("Real Madrid Hats")
        self.assertEqual(anchored["rule_type"], "Title + Type")
        self.assertIn("Product type equals 'Hats'", anchored["suggested_rule"])
        broad = self.classify_one("Real Madrid Jerseys & Gear")
        self.assertEqual(broad["rule_type"], "Club (metafield)")
        self.assertNotIn("Product type", broad["suggested_rule"])

    def test_whole_type_threshold(self):
        ok = self.classify_one("Posters", count=5)
        self.assertEqual((ok["classification"], ok["rule_type"]), ("AUTOMATABLE", "Product type"))
        small = self.classify_one("Posters", count=2)
        self.assertEqual(small["classification"], "REVIEW")
        self.assertIn("curated subset", small["suggested_rule"])

    def test_keep_manual_and_ambiguous(self):
        self.assertEqual(self.classify_one("World Cup 2026")["classification"], "KEEP_MANUAL")
        self.assertEqual(self.classify_one("Santos Jerseys")["classification"], "REVIEW")

    def test_adds_math(self):
        members = {"gid://shopify/Collection/1": {"gid://shopify/Product/20"}}
        row = self.classify_one("Real Madrid Jerseys & Gear", count=1, members=members)
        self.assertEqual(row["adds"], 2)
        self.assertEqual(row["count_after"], 3)

    def test_members_loader(self):
        import tempfile
        lines = ['{"id": "gid://shopify/Collection/9"}',
                 '{"id": "gid://shopify/Product/1", "__parentId": "gid://shopify/Collection/9"}',
                 '{"id": "gid://shopify/Product/2", "__parentId": "gid://shopify/Collection/9"}']
        with tempfile.NamedTemporaryFile("w", suffix=".jsonl", delete=False,
                                         encoding="utf-8") as fh:
            fh.write("\n".join(lines))
            path = fh.name
        try:
            members = ca.load_members(path)
        finally:
            os.unlink(path)
        self.assertEqual(len(members["gid://shopify/Collection/9"]), 2)


class TestDeliverable(unittest.TestCase):
    def test_xlsx_and_split(self):
        try:
            import openpyxl
        except ImportError:
            raise unittest.SkipTest("openpyxl not installed")
        import tempfile
        import deliverable as dlv
        rows = [
            {"title": "Real Madrid Hats", "handle": "rm-hats", "count": 3,
             "classification": "AUTOMATABLE", "rule_type": "Title + Type",
             "suggested_rule": "x", "adds": 2, "count_after": 5,
             "admin_url": "https://admin.shopify.com/store/s/collections/151515151515"},
            {"title": "Blue Cleats", "handle": "bc", "count": 4,
             "classification": "COLOUR_VERIFY", "rule_type": "Title colour",
             "suggested_rule": "y", "adds": "", "count_after": "",
             "admin_url": "https://admin.shopify.com/store/s/collections/2"},
            {"title": "Weird", "handle": "w", "count": 1, "classification": "REVIEW",
             "rule_type": "Review", "suggested_rule": "z", "adds": "", "count_after": "",
             "admin_url": "https://admin.shopify.com/store/s/collections/3"},
            {"title": "Sale", "handle": "sale", "count": 9, "classification": "KEEP_MANUAL",
             "rule_type": "", "suggested_rule": "", "adds": "", "count_after": "",
             "admin_url": "https://admin.shopify.com/store/s/collections/4"},
        ]
        with tempfile.TemporaryDirectory() as td:
            import config as cfgmod
            old = cfgmod.OUT_DIR
            cfgmod.OUT_DIR = td
            try:
                path = dlv.build_xlsx(rows, "test", log=lambda *a: None)
                wb = openpyxl.load_workbook(path)
            finally:
                cfgmod.OUT_DIR = old
        self.assertEqual(wb.sheetnames, ["Ready to Automate", "Needs Review"])
        ws = wb["Ready to Automate"]
        self.assertEqual(ws.max_row, 3)
        self.assertEqual(ws.cell(row=2, column=1).number_format, "@")
        self.assertEqual(ws.cell(row=2, column=1).value, "151515151515")
        self.assertEqual(wb["Needs Review"].max_row, 2)
if __name__ == "__main__":
    unittest.main()
